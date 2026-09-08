"""
Presentation generation (BRD 2.8): editable PowerPoint + PDF, plus the
buyer credit-limit table as an editable Excel file (BRD 2.6 output).

Page order follows the BRD 2.8 table exactly:
  1. Client cover        client name, date, new-business/renewal title
  2. About the broker    fixed template text
  3. Important info      fixed regulatory wording + insurers approached +
                         auto-generated declined line
  4. Terms comparison    the 16 BRD 2.3 rows x 1-6 insurer columns, with
                         the recommended column highlighted, free-format
                         notes beneath
  5. Buyer credit limits omitted cleanly when no limits were supplied
  6. Comments & rec      standard wording with the insurer's name merged,
                         plus the broker's reasons
  7. Contact             fixed template text

Rules enforced here, server-side:
  - BRD 2.5 gate: generation refused (409) until the four key values are
    confirmed — no client can bypass it.
  - BRD 2.7: the regulatory wording is always rendered; it cannot be
    edited out of a generated presentation.
  - BRD 2.2: blank values render blank — never "N/A", never a default.
  - The PPTX uses only standard shapes/tables/text so it opens and edits
    cleanly in Google Slides (the brokerage runs on G Suite).

The visual design is a neutral implementation of the brokerage palette;
the client's approved template (BRD open item) can restyle it without
touching the data flow.
"""

import io
import re
from datetime import date

import openpyxl
import pymupdf
from openpyxl.styles import Font as XlsxFont
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.core.errors import ExportBlockedError, InvalidDocumentError
from app.models.presentation import PRESENTATION_ROWS, PresentationRequest
from app.models.schemas import CONFIRM_REQUIRED_FIELDS

# ── Palette (wireframe tokens) ───────────────────────────────────────────
NAVY = (0x14, 0x1B, 0x3D)
INK = (0x0F, 0x17, 0x29)
INK2 = (0x55, 0x60, 0x72)
INK3 = (0x9A, 0xA4, 0xB2)
ACCENT = (0x4F, 0x46, 0xE5)
REC_FILL = (0xEE, 0xF0, 0xFE)
PANEL = (0xF4, 0xF5, 0xF8)
WHITE = (0xFF, 0xFF, 0xFF)
MUTED_ON_NAVY = (0x8F, 0xA2, 0xC9)

# ── Fixed template wording (BRD 2.7/2.8 — never editable out) ────────────
ABOUT_TEXT = (
    "UK Credit Insurance Brokers is an independent, specialist credit "
    "insurance brokerage. We arrange and manage trade credit insurance for "
    "UK businesses, working with the leading insurers in the market to "
    "protect our clients against non-payment and insolvency risk. Our team "
    "handles placement, credit-limit management and claims support "
    "throughout the life of every policy."
)
REGULATORY_TEXT = (
    "UK Credit Insurance Brokers is authorised and regulated by the "
    "Financial Conduct Authority. This document is a summary of the "
    "quotations obtained on your behalf and does not amend the policy "
    "documents; in all cases the policy wording prevails. Premiums are "
    "shown exclusive of Insurance Premium Tax unless stated otherwise. "
    "This document is intended solely for the client named on the front "
    "page and should not be shared with any third party."
)
RECOMMENDATION_WORDING = (
    "Having reviewed the quotations obtained on your behalf, we recommend "
    "{name}. This recommendation reflects the cover, terms and pricing "
    "offered relative to the alternatives presented. This document is a "
    "summary; the policy wording prevails. UK Credit Insurance Brokers is "
    "authorised and regulated by the Financial Conduct Authority."
)
CONTACT_LINES = [
    "UK Credit Insurance Brokers",
    "Underwriting desk",
    "enquiries@ukcib.co.uk",
    "020 7000 0000",
]


def cover_title(req: PresentationRequest) -> str:
    """BRD: one template, two titles — the front page is the only difference."""
    return (
        "Renewal of Credit Insurance"
        if req.project_type == "renewal"
        else "Credit Insurance Proposals"
    )


def check_export_gate(req: PresentationRequest) -> None:
    """BRD 2.5/2.8: export blocked until the four key values are confirmed."""
    missing = [f for f in CONFIRM_REQUIRED_FIELDS if f not in req.confirmed_fields]
    if missing:
        raise ExportBlockedError(
            "Export is blocked until estimated annual premium, indemnity, "
            "excess and max annual liability are confirmed on the review "
            f"screen ({len(missing)} remaining)."
        )


def declined_insurers(req: PresentationRequest) -> list[str]:
    """
    BRD 2.1/S8: approached insurers with no quote column are auto-named as
    declined; a late quote (new column) moves them off this line.
    """
    column_names = [c.name.lower() for c in req.columns]
    declined = []
    for name in req.approached_insurers:
        needle = name.lower()
        quoted = any(needle in cn or cn in needle for cn in column_names)
        if not quoted:
            declined.append(name)
    return declined


# ── Shared content assembly — one source of truth for both renderers ────

def _kicker(req: PresentationRequest) -> str:
    return "RENEWAL" if req.project_type == "renewal" else "NEW BUSINESS"


def _subtitle(req: PresentationRequest) -> str:
    subtitle = f"{req.client_name} · {date.today().strftime('%B %Y')}"
    return f"{subtitle} · {req.reference}" if req.reference else subtitle


def _quoted_sentence(req: PresentationRequest) -> str:
    names = [c.name for c in req.columns]
    plural = "s" if len(names) != 1 else ""
    return f"{', '.join(names)} — quotation{plural} obtained."


def _declined_sentence(req: PresentationRequest) -> str | None:
    declined = declined_insurers(req)
    if not declined:
        return None
    verb = "was" if len(declined) == 1 else "were"
    return f"{', '.join(declined)} {verb} approached but declined to quote."


def _recommended_name(req: PresentationRequest) -> str:
    col = next((c for c in req.columns if _is_rec(req, c)), None)
    return col.name if col else "[no insurer selected]"


def _is_rec(req: PresentationRequest, col) -> bool:
    return col.id == req.recommended_id


def _limits_headers(req: PresentationRequest) -> list[str]:
    return ["Buyer", "Company no.", "Required"] + [c.name for c in req.columns]


def suggested_filename(req: PresentationRequest, extension: str) -> str:
    # ASCII-only: HTTP headers are latin-1 and a non-ASCII client name must
    # never be able to break the download response.
    client = re.sub(r"[^A-Za-z0-9 \-]", "", req.client_name).strip() or "Client"
    return f"{client[:80]} - {cover_title(req)}.{extension}"


def _cell_text(value: str | None) -> str:
    # BRD 2.2: blank stays blank — never a placeholder.
    return (value or "").strip()


# ═════════════════════════════ PPTX ═════════════════════════════════════

def _rgb(t: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*t)


def _add_textbox(slide, x, y, w, h, text, *, size, color=INK, bold=False,
                 align=None, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = wrap
    paragraph = frame.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _fill_slide(slide, color: tuple[int, int, int]) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _style_cell(cell, text, *, size, bold=False, color=INK, fill=None):
    cell.text_frame.word_wrap = True
    paragraph = cell.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
    else:
        cell.fill.background()


def build_pptx(req: PresentationRequest) -> bytes:
    """Render the full deck as an editable, Google-Slides-clean .pptx."""
    check_export_gate(req)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ── 1. Client cover ──────────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _fill_slide(slide, NAVY)
    _add_textbox(slide, 0.9, 2.2, 8, 0.5, _kicker(req), size=13, color=MUTED_ON_NAVY, bold=True)
    _add_textbox(slide, 0.9, 2.7, 11.5, 1.6, cover_title(req), size=44, color=WHITE, bold=True)
    _add_textbox(slide, 0.9, 4.35, 11.5, 0.6, _subtitle(req), size=16, color=(0xC3, 0xCF, 0xE0))

    # ── 2. About the broker (fixed) ──────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.9, 0.7, 11.5, 0.7, "About the broker", size=28, bold=True)
    _add_textbox(slide, 0.9, 1.7, 11.5, 4.5, ABOUT_TEXT, size=14, color=INK2)

    # ── 3. Important information ─────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.9, 0.7, 11.5, 0.7, "Important information", size=28, bold=True)
    _add_textbox(slide, 0.9, 1.6, 11.5, 2.2, REGULATORY_TEXT, size=12, color=INK2)
    _add_textbox(slide, 0.9, 3.9, 11.5, 0.4, "INSURERS APPROACHED", size=11,
                 color=INK3, bold=True)
    _add_textbox(slide, 0.9, 4.35, 11.5, 0.5, _quoted_sentence(req), size=12)
    declined_line = _declined_sentence(req)
    if declined_line:
        _add_textbox(slide, 0.9, 4.9, 11.5, 0.5, declined_line,
                     size=12, color=(0xB4, 0x53, 0x09))

    # ── 4. Terms comparison ──────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.6, 0.35, 12, 0.55, "Terms comparison", size=24, bold=True)
    n_cols = len(req.columns)
    n_rows = 1 + len(PRESENTATION_ROWS)
    table_w = 12.1
    label_w = 2.7
    value_w = (table_w - label_w) / n_cols
    shape = slide.shapes.add_table(
        n_rows, 1 + n_cols, Inches(0.6), Inches(1.0), Inches(table_w), Inches(5.7)
    )
    table = shape.table
    table.columns[0].width = Inches(label_w)
    for i in range(n_cols):
        table.columns[i + 1].width = Inches(value_w)

    _style_cell(table.cell(0, 0), "Field", size=10, bold=True, color=INK3, fill=PANEL)
    for c, col in enumerate(req.columns, start=1):
        rec = _is_rec(req, col)
        _style_cell(table.cell(0, c), col.name, size=10.5, bold=True,
                    color=ACCENT if rec else INK, fill=REC_FILL if rec else PANEL)
    for r, (key, label) in enumerate(PRESENTATION_ROWS, start=1):
        _style_cell(table.cell(r, 0), label, size=9, bold=True, color=INK2, fill=None)
        for c, col in enumerate(req.columns, start=1):
            rec = _is_rec(req, col)
            _style_cell(table.cell(r, c), _cell_text(col.values.get(key)),
                        size=9, fill=REC_FILL if rec else None)
    if req.notes.strip():
        _add_textbox(slide, 0.6, 6.85, 12.1, 0.55, req.notes.strip(),
                     size=9, color=INK2)

    # ── 5. Buyer credit limits (omitted cleanly when none) ───────────────
    if req.credit_limits:
        slide = prs.slides.add_slide(blank)
        _add_textbox(slide, 0.6, 0.5, 12, 0.6, "Buyer credit limits", size=24, bold=True)
        rows = 1 + len(req.credit_limits)
        cols = 3 + n_cols
        shape = slide.shapes.add_table(
            rows, cols, Inches(0.6), Inches(1.3), Inches(12.1),
            Inches(min(5.6, 0.4 * rows)),
        )
        table = shape.table
        headers = _limits_headers(req)
        for c, header in enumerate(headers):
            col_obj = req.columns[c - 3] if c >= 3 else None
            rec = col_obj is not None and col_obj.id == req.recommended_id
            _style_cell(table.cell(0, c), header, size=10, bold=True,
                        color=ACCENT if rec else INK3, fill=REC_FILL if rec else PANEL)
        for r, row in enumerate(req.credit_limits, start=1):
            _style_cell(table.cell(r, 0), _cell_text(row.buyer), size=9.5)
            _style_cell(table.cell(r, 1), _cell_text(row.company_number), size=9.5)
            _style_cell(table.cell(r, 2), _cell_text(row.required), size=9.5)
            for c, col in enumerate(req.columns, start=3):
                rec = _is_rec(req, col)
                _style_cell(table.cell(r, c), _cell_text(row.offers.get(col.id)),
                            size=9.5, fill=REC_FILL if rec else None)

    # ── 6. Comments and recommendation ───────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _add_textbox(slide, 0.9, 0.7, 11.5, 0.7, "Comments and recommendation",
                 size=28, bold=True)
    _add_textbox(slide, 0.9, 1.7, 11.5, 1.8,
                 RECOMMENDATION_WORDING.format(name=_recommended_name(req)),
                 size=13, color=INK2)
    if req.reasons.strip():
        _add_textbox(slide, 0.9, 3.7, 11.5, 0.4, "REASONS FOR THE RECOMMENDATION",
                     size=11, color=INK3, bold=True)
        _add_textbox(slide, 0.9, 4.15, 11.5, 2.6, req.reasons.strip(),
                     size=12.5, color=INK)

    # ── 7. Contact (fixed) ───────────────────────────────────────────────
    slide = prs.slides.add_slide(blank)
    _fill_slide(slide, NAVY)
    _add_textbox(slide, 0.9, 2.5, 11.5, 0.7, "Contact", size=28, color=WHITE, bold=True)
    _add_textbox(slide, 0.9, 3.4, 11.5, 1.8, "\n".join(CONTACT_LINES),
                 size=14, color=(0xC3, 0xCF, 0xE0))

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ═════════════════════════════ PDF ══════════════════════════════════════

PAGE_W, PAGE_H = 960, 540  # 16:9 points
MARGIN = 54


def _norm(color: tuple[int, int, int]) -> tuple[float, float, float]:
    return tuple(v / 255 for v in color)


_PDF_CHAR_MAP = str.maketrans({
    "—": "-", "–": "-",          # em/en dash
    "‘": "'", "’": "'",          # curly single quotes
    "“": '"', "”": '"',          # curly double quotes
    "…": "...",
})


def _pdf_safe(text: str) -> str:
    """Base-14 Helvetica is Latin-1 only — map common typographic chars."""
    return text.translate(_PDF_CHAR_MAP)


def _pdf_text(page, x, y, w, h, text, *, size, color=INK, bold=False):
    # insert_textbox silently drops text whose rect is shorter than the
    # rendered line height — guarantee headroom so titles can never vanish.
    h = max(h, size * 2)
    page.insert_textbox(
        pymupdf.Rect(x, y, x + w, y + h), _pdf_safe(text),
        fontsize=size, fontname="hebo" if bold else "helv",
        color=_norm(color),
    )


def _pdf_table(page, x, y, widths, cells, row_h, *, font_size):
    """cells: rows of (text, fill|None, bold, color) tuples."""
    for r, row in enumerate(cells):
        cx = x
        for c, (text, fill, bold, color) in enumerate(row):
            rect = pymupdf.Rect(cx, y + r * row_h, cx + widths[c], y + (r + 1) * row_h)
            if fill is not None:
                page.draw_rect(rect, color=None, fill=_norm(fill))
            page.draw_rect(rect, color=_norm((0xE8, 0xEB, 0xF0)), width=0.5)
            page.insert_textbox(
                rect + (4, 3, -4, -1), _pdf_safe(text), fontsize=font_size,
                fontname="hebo" if bold else "helv", color=_norm(color),
            )
            cx += widths[c]


def build_pdf(req: PresentationRequest) -> bytes:
    """Render the same page sequence as the PPTX, as a PDF."""
    check_export_gate(req)
    doc = pymupdf.open()

    # ── 1. Cover ─────────────────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.draw_rect(page.rect, color=None, fill=_norm(NAVY))
    _pdf_text(page, MARGIN, 170, 500, 24, _kicker(req), size=11, color=MUTED_ON_NAVY, bold=True)
    _pdf_text(page, MARGIN, 200, 840, 60, cover_title(req), size=40, color=WHITE, bold=True)
    _pdf_text(page, MARGIN, 280, 840, 26, _subtitle(req), size=14, color=(0xC3, 0xCF, 0xE0))

    # ── 2. About the broker ──────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 50, 840, 34, "About the broker", size=24, bold=True)
    _pdf_text(page, MARGIN, 110, 840, 300, ABOUT_TEXT, size=12, color=INK2)

    # ── 3. Important information ─────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 50, 840, 34, "Important information", size=24, bold=True)
    _pdf_text(page, MARGIN, 105, 840, 150, REGULATORY_TEXT, size=10.5, color=INK2)
    _pdf_text(page, MARGIN, 280, 840, 18, "INSURERS APPROACHED", size=9,
              color=INK3, bold=True)
    _pdf_text(page, MARGIN, 302, 840, 22, _quoted_sentence(req), size=11)
    declined_line = _declined_sentence(req)
    if declined_line:
        _pdf_text(page, MARGIN, 328, 840, 40, declined_line,
                  size=11, color=(0xB4, 0x53, 0x09))

    # ── 4. Terms comparison ──────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 30, 840, 30, "Terms comparison", size=20, bold=True)
    n_cols = len(req.columns)
    label_w = 190.0
    value_w = (PAGE_W - 2 * MARGIN - label_w) / n_cols
    widths = [label_w] + [value_w] * n_cols
    header = [("Field", PANEL, True, INK3)] + [
        (c.name, REC_FILL if _is_rec(req, c) else PANEL, True,
         ACCENT if _is_rec(req, c) else INK)
        for c in req.columns
    ]
    body = []
    for key, label in PRESENTATION_ROWS:
        row = [(label, None, True, INK2)]
        for col in req.columns:
            rec = _is_rec(req, col)
            row.append((_cell_text(col.values.get(key)), REC_FILL if rec else None,
                        False, INK))
        body.append(row)
    _pdf_table(page, MARGIN, 68, widths, [header] + body, 26, font_size=8)
    if req.notes.strip():
        _pdf_text(page, MARGIN, 68 + 26 * (1 + len(PRESENTATION_ROWS)) + 8,
                  840, 30, req.notes.strip(), size=8, color=INK2)

    # ── 5. Buyer credit limits (omitted cleanly when none) ───────────────
    if req.credit_limits:
        page = doc.new_page(width=PAGE_W, height=PAGE_H)
        _pdf_text(page, MARGIN, 40, 840, 30, "Buyer credit limits", size=20, bold=True)
        fixed = [170.0, 110.0, 100.0]
        value_w = (PAGE_W - 2 * MARGIN - sum(fixed)) / n_cols
        widths = fixed + [value_w] * n_cols
        header = [("Buyer", PANEL, True, INK3), ("Company no.", PANEL, True, INK3),
                  ("Required", PANEL, True, INK3)] + [
            (c.name, REC_FILL if _is_rec(req, c) else PANEL, True,
             ACCENT if _is_rec(req, c) else INK)
            for c in req.columns
        ]
        body = []
        for row in req.credit_limits:
            cells = [(_cell_text(row.buyer), None, False, INK),
                     (_cell_text(row.company_number), None, False, INK2),
                     (_cell_text(row.required), None, False, INK)]
            for col in req.columns:
                rec = _is_rec(req, col)
                cells.append((_cell_text(row.offers.get(col.id)),
                              REC_FILL if rec else None, False, INK))
            body.append(cells)
        _pdf_table(page, MARGIN, 82, widths, [header] + body, 28, font_size=9)

    # ── 6. Comments and recommendation ───────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    _pdf_text(page, MARGIN, 50, 840, 34, "Comments and recommendation",
              size=24, bold=True)
    _pdf_text(page, MARGIN, 110, 840, 120,
              RECOMMENDATION_WORDING.format(name=_recommended_name(req)),
              size=12, color=INK2)
    if req.reasons.strip():
        _pdf_text(page, MARGIN, 250, 840, 18, "REASONS FOR THE RECOMMENDATION",
                  size=9, color=INK3, bold=True)
        _pdf_text(page, MARGIN, 272, 840, 200, req.reasons.strip(), size=11)

    # ── 7. Contact ───────────────────────────────────────────────────────
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.draw_rect(page.rect, color=None, fill=_norm(NAVY))
    _pdf_text(page, MARGIN, 190, 840, 34, "Contact", size=24, color=WHITE, bold=True)
    _pdf_text(page, MARGIN, 240, 840, 120, "\n".join(CONTACT_LINES),
              size=13, color=(0xC3, 0xCF, 0xE0))

    data = doc.tobytes()
    doc.close()
    return data


# ═════════════════════════════ XLSX ═════════════════════════════════════

def build_limits_xlsx(req: PresentationRequest) -> bytes:
    """BRD 2.6 output: the buyer credit-limit table as an editable Excel file."""
    check_export_gate(req)
    if not req.credit_limits:
        raise InvalidDocumentError(
            "No buyer credit limits to export — the credit-limit page is "
            "omitted when no limits are supplied."
        )
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Buyer credit limits"
    headers = _limits_headers(req)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = XlsxFont(bold=True)
    for row in req.credit_limits:
        sheet.append(
            [row.buyer, row.company_number, row.required]
            + [row.offers.get(c.id, "") for c in req.columns]
        )
    for i, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 4)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
