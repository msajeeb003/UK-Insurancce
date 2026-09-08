"""Presentation generation (BRD 2.8) — gate, page set, formats, declined line."""

import io

import openpyxl
import pymupdf
import pytest
from pptx import Presentation

from app.core.errors import ExportBlockedError, InvalidDocumentError
from app.models.presentation import (
    CreditLimitRow,
    PresentationColumn,
    PresentationRequest,
)
from app.models.schemas import CONFIRM_REQUIRED_FIELDS
from app.services.presentation import (
    build_limits_xlsx,
    build_pdf,
    build_pptx,
    declined_insurers,
)

ALL_CONFIRMED = list(CONFIRM_REQUIRED_FIELDS)


def make_request(**overrides) -> PresentationRequest:
    base = dict(
        client_name="Aldgate Timber Ltd",
        reference="UKCIB-2418",
        project_type="new",
        columns=[
            PresentationColumn(id="a", name="Allianz Trade", values={
                "type": "Whole Turnover", "annual_turnover": "£4,500,000",
                "premium_rate": "0.32%", "estimated_annual_premium_exc_ipt": "£14,400",
                "indemnity": "90%", "excess": "£1,000",
                "excess_type": "Minimum Retention", "debt": "Included",
            }),
            PresentationColumn(id="b", name="Atradius", values={
                "type": "Whole Turnover", "annual_turnover": "£4,500,000",
                "premium_rate": "0.28%", "estimated_annual_premium_exc_ipt": "£12,600",
                "indemnity": "90%", "excess": "£2,500",
                "excess_type": "Deductible", "debt": "Included",
            }),
        ],
        recommended_id="b",
        approached_insurers=["Allianz Trade", "Atradius", "Coface"],
        credit_limits=[
            CreditLimitRow(buyer="Meridian Foods Ltd", company_number="04821990",
                           required="£250,000", offers={"a": "£250,000", "b": "£200,000"}),
        ],
        notes="All quotes exclude IPT.",
        reasons="Best combination of rate and cover.",
        confirmed_fields=ALL_CONFIRMED,
    )
    base.update(overrides)
    return PresentationRequest(**base)


# ── Gate (BRD 2.5/2.8) ───────────────────────────────────────────────────

def test_gate_blocks_unconfirmed_export():
    req = make_request(confirmed_fields=["indemnity"])
    with pytest.raises(ExportBlockedError, match="3 remaining"):
        build_pptx(req)
    with pytest.raises(ExportBlockedError):
        build_pdf(req)


# ── Declined line (BRD 2.1/S8) ───────────────────────────────────────────

def test_declined_is_approached_minus_quoted():
    req = make_request()
    assert declined_insurers(req) == ["Coface"]


def test_late_quote_moves_off_the_declined_line():
    req = make_request()
    req.columns.append(PresentationColumn(id="c", name="Coface", values={}))
    assert declined_insurers(req) == []


# ── PPTX ─────────────────────────────────────────────────────────────────

def test_pptx_has_the_brd_page_set():
    deck = Presentation(io.BytesIO(build_pptx(make_request())))
    assert len(deck.slides) == 7  # cover, about, info, terms, limits, rec, contact

    all_text = "\n".join(
        shape.text_frame.text
        for slide in deck.slides for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Credit Insurance Proposals" in all_text
    assert "Aldgate Timber Ltd" in all_text
    # Regulatory wording is always present (BRD 2.7).
    assert "Financial Conduct Authority" in all_text
    # Declined auto-line and the merged recommendation name.
    assert "Coface" in all_text and "declined to quote" in all_text
    assert "we recommend Atradius" in all_text
    # A blank value never becomes a placeholder.
    assert "N/A" not in all_text


def test_pptx_renewal_title_and_omitted_limits():
    req = make_request(project_type="renewal", credit_limits=[])
    deck = Presentation(io.BytesIO(build_pptx(req)))
    assert len(deck.slides) == 6  # credit-limit page omitted cleanly
    covers = deck.slides[0]
    text = "\n".join(s.text_frame.text for s in covers.shapes if s.has_text_frame)
    assert "Renewal of Credit Insurance" in text


def test_pptx_terms_table_holds_single_and_six_columns():
    for n in (1, 6):
        cols = [PresentationColumn(id=f"c{i}", name=f"Insurer {i}", values={})
                for i in range(n)]
        deck = Presentation(io.BytesIO(build_pptx(make_request(
            columns=cols, recommended_id="c0", credit_limits=[]))))
        terms = deck.slides[3]
        table = next(s for s in terms.shapes if s.has_table).table
        assert len(table.columns) == 1 + n
        assert len(table.rows) == 1 + 15  # header + the BRD 2.3 rows


# ── PDF ──────────────────────────────────────────────────────────────────

def test_pdf_mirrors_the_page_set():
    doc = pymupdf.open(stream=build_pdf(make_request()), filetype="pdf")
    assert doc.page_count == 7
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    assert "Terms comparison" in text
    assert "we recommend Atradius" in text
    assert "Financial Conduct Authority" in text
    assert "Meridian Foods Ltd" in text


def test_pdf_omits_limits_page_when_none():
    doc = pymupdf.open(
        stream=build_pdf(make_request(credit_limits=[])), filetype="pdf"
    )
    assert doc.page_count == 6
    doc.close()


# ── Credit-limits Excel (BRD 2.6 output) ─────────────────────────────────

def test_limits_xlsx_round_trips():
    workbook = openpyxl.load_workbook(io.BytesIO(build_limits_xlsx(make_request())))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Buyer", "Company no.", "Required", "Allianz Trade", "Atradius")
    assert rows[1] == ("Meridian Foods Ltd", "04821990", "£250,000", "£250,000", "£200,000")


def test_limits_xlsx_refused_when_empty():
    with pytest.raises(InvalidDocumentError):
        build_limits_xlsx(make_request(credit_limits=[]))


# ── API endpoint ─────────────────────────────────────────────────────────

def test_endpoint_gate_and_download_headers(client):
    payload = make_request().model_dump()

    blocked = dict(payload, confirmed_fields=[])
    res = client.post("/generate-presentation?format=pptx", json=blocked)
    assert res.status_code == 409

    res = client.post("/generate-presentation?format=pptx", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml"
    )
    assert 'filename="Aldgate Timber Ltd - Credit Insurance Proposals.pptx"' in (
        res.headers["content-disposition"]
    )
    assert Presentation(io.BytesIO(res.content))  # valid, editable pptx

    res = client.post("/generate-presentation?format=pdf", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
